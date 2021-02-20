import sys
import os
from urllib.parse import urlparse
import boto3
import time
from tdp import DocumentProcessor
from og import OutputGenerator
from helper import FileHelper, S3Helper
import boto3
from botocore.config import Config
import boto3
import shutil
import pandas as pd
import time
import re
import shutil
import os
import math
from tika import parser
import openpyxl
import psycopg2
import datetime

def move_po(data2):
    shutil.copy(data2,'/home/ubuntu/Desktop/Automation/FINAL_CODES/FINAL_CODES')

def extract(pdf_file,o):
    class Textractor:
        def getInputParameters(self, args):
            event = {}
            i = 0
            print(args)
            args=args.split()
            
            if(args):
                while(i < len(args)):
                    if(args[i] == '--documents'):
                        event['documents'] = args[i+1]
                        i = i + 1
                    if(args[i] == '--region'):
                        event['region'] = args[i+1]
                        i = i + 1
                    if(args[i] == '--text'):
                        event['text'] = True
                    if(args[i] == '--forms'):
                        event['forms'] = True
                    if(args[i] == '--tables'):
                        event['tables'] = True
                    if(args[i] == '--insights'):
                        event['insights'] = True
                    if(args[i] == '--medical-insights'):
                        event['medical-insights'] = True
                    if(args[i] == '--translate'):
                        event['translate'] = args[i+1]
                        i = i + 1

                    i = i + 1
            print(event)
            return event

        def validateInput(self, args):

            event = self.getInputParameters(args)

            ips = {}

            if(not 'documents' in event):
                raise Exception("Document or path to a foler or S3 bucket containing documents is required.")

            inputDocument = event['documents']
            idl = inputDocument.lower()

            bucketName = None
            documents = []
            awsRegion = 'us-east-1'

            if(idl.startswith("s3://")):
                o = urlparse(inputDocument)
                bucketName = o.netloc
                path = o.path[1:]
                ar = S3Helper.getS3BucketRegion(bucketName)
                if(ar):
                    awsRegion = ar

                if(idl.endswith("/")):
                    allowedFileTypes = ["jpg", "jpeg", "png", "pdf"]
                    documents = S3Helper.getFileNames(awsRegion, bucketName, path, 1, allowedFileTypes)
                else:
                    documents.append(path)
            else:
                if(idl.endswith("/")):
                    allowedFileTypes = ["jpg", "jpeg", "png"]
                    documents = FileHelper.getFileNames(inputDocument, allowedFileTypes)
                else:
                    documents.append(inputDocument)

                if('region' in event):
                    awsRegion = event['region']

            ips["bucketName"] = bucketName
            ips["documents"] = documents
            ips["awsRegion"] = awsRegion
            ips["text"] = ('text' in event)
            ips["forms"] = ('forms' in event)
            ips["tables"] = ('tables' in event)
            ips["insights"] = ('insights' in event)
            ips["medical-insights"] = ('medical-insights' in event)
            if("translate" in event):
                ips["translate"] = event["translate"]
            else:
                ips["translate"] = ""

            return ips

        def processDocument(self, ips, i, document):
            print("\nTextracting Document # {}: {}".format(i, document))
            print('=' * (len(document)+30))

            # Get document textracted
            dp = DocumentProcessor(ips["bucketName"], document, ips["awsRegion"], ips["text"], ips["forms"], ips["tables"])
            response = dp.run()

            if(response):
                print("Recieved Textract response...")
                #FileHelper.writeToFile("temp-response.json", json.dumps(response))

                #Generate output files
                print("Generating output...")
                name, ext = FileHelper.getFileNameAndExtension(document)
                opg = OutputGenerator(response,
                            "{}-{}".format(name, ext),
                            ips["forms"], ips["tables"])
                opg.run()

                if(ips["insights"] or ips["medical-insights"] or ips["translate"]):
                    opg.generateInsights(ips["insights"], ips["medical-insights"], ips["translate"], ips["awsRegion"])

                print("{} textracted successfully.".format(document))
            else:
                print("Could not generate output for {}.".format(document))

        def printFormatException(self, e):
            print("Invalid input: {}".format(e))
            print("Valid format:")
            print('- python3 textractor.py --documents mydoc.jpg --text --forms --tables --region us-east-1')
            print('- python3 textractor.py --documents ./myfolder/ --text --forms --tables')
            print('- python3 textractor.py --documents s3://mybucket/mydoc.pdf --text --forms --tables')
            print('- python3 textractor.py --documents s3://mybucket/ --text --forms --tables')

        def run(self,argv):

            ips = None
            try:
                ips = self.validateInput(argv)
            except Exception as e:
                self.printFormatException(e)

            #try:
            i = 1
            totalDocuments = len(ips["documents"])

            print("\n")
            print('*' * 60)
            print("Total input documents: {}".format(totalDocuments))
            print('*' * 60)

            for document in ips["documents"]:
                self.processDocument(ips, i, document)

                remaining = len(ips["documents"])-i

                if(remaining > 0):
                    print("\nRemaining documents: {}".format(remaining))

                    # print("\nTaking a short break...")
                    # time.sleep(20)
                    # print("Allright, ready to go...\n")

                i = i + 1

            print("\n")
            print('*' * 60)
            print("Successfully textracted documents: {}".format(totalDocuments))
            print('*' * 60)
            print("\n")
            #except Exception as e:
            #    print("Something went wrong:\n====================================================\n{}".format(e))

    #Textractor().run()


    def upload_cloud(pdf,name):
        s3 = boto3.resource('s3')
        BUCKET = "testpdfs1200"
        s3.Bucket(BUCKET).upload_file(pdf, name)
        print("uploaded")




    input_pdf_path=r"/home/ubuntu/Desktop/Automation/FINAL_CODES/FINAL_CODES/FETCHED_PDF"
    for data in os.listdir(input_pdf_path):
           data2= (input_pdf_path+"/%s")%data
           if ".pdf" in data2:
               name=data.split("/")[-1]
               pdf=data2
               print(pdf,"----",name)
               move_po(data2)
               break
##    upload_cloud(pdf,name)               

               

    def api_request():
        s=Textractor()
        #s.getInputParameters(r'--documents s3://testpdfs1200/AKNA__PO_2112_GGN47.pdf --text --forms --tables')
        print(f'--documents s3://testpdfs1200/{name}')

        s.validateInput(f'--documents s3://testpdfs1200//{name} --tables')
        s.run(f'--documents s3://testpdfs1200/{name} --tables')
##    api_request()

    wb=openpyxl.Workbook()
    sh1=wb.active
    header=["vendor_name", "po_no", "po_date","des","qty","uom","mrp","foc","prate","disc_percent","disc_amount","total_disc_percent","total_disc_amount","sgst","igst","cgst","pdf_file","time_stamp"]
    sh1.append(header)
    wb.save("excel.xlsx")


    def combine_csv(csv_path1,csv_path2):
        df1=pd.read_csv(csv_path1)
        try:
            df2=pd.read_csv(csv_path2)
            df3=[df1,df2]
            df3=pd.concat(df3)
        except:
            df3=df1
        df3.to_csv(r'final.csv')

    def checker(quantity,UMRP,TUMRP):
        if float(quantity)*float(UMRP)==float(TUMRP):
            return UMRP,TUMRP
        return "",""



    def checker2(main_data):
        print("main_data",main_data)
        if re.search("\s+(EACH|BOTTLE|UNIT).*(EACH|BOTTLE|UNIT).*(EACH|BOTTLE|UNIT).*(EACH|BOTTLE|UNIT)",main_data):
            return "Tika"
        else:
            return ""
    def extract_line_item(data,indx1,indx2,indx3):
        #des=data[indx1[0]]
        #quantity=data[indx2[0]].split()[indx2[1]]
        #uom=data[indx3[0]].split()[indx3[1]]
        rate=data[indx1]
        UMRP=str(data[indx2])
        TUMRP=str(data[indx3])
        #print("final data",des,rate,uom,quantity,UMRP,TUMRP)
        #print("main data",des,rate,uom,quantity,UMRP,TUMRP)
        return rate,UMRP,TUMRP

    def display(des,rate,uom,quantity,UMRP,TUMRP,discount_per):
        vendor_name=''
        print("des=",des)
        print("rate=",rate)
        print("uom=",uom)
        print("quantity=",quantity)
        print("UMRP=",UMRP)
        print("TUMRP=",TUMRP)
        print("discount",discount_per)
        print()
        print("-----------------------------------------------------------------------------------------------------------------------------------------")


    def count_line(path):
        count_line_item=0
        df=pd.read_csv("final.csv")
        for i in range(df.shape[0]):
            des,rate,uom,quantity,UMRP,TUMRP="","","","","",""
            data=df.iloc[i,::].values.tolist()
            for n,y in enumerate(data):
                data[n]=str(y)
            #print("---------------------------------------------------------",len(data),data)
            if 'Qty UOM' not in data and 'Purchase ' not in data and 'Rate ' not in data and 'Discounted Value ' not in data and len(data)>=6 and data.count("nan")<4:
                if "PO Date " in data:
                    break
                for n,y in enumerate(data):
                    data[n]=str(y)
                count_line_item=count_line_item+1
                #print("count_line_item",count_line_item)
        return count_line_item
                
    def extract_from_tika(pdf,check):
        conn=psycopg2.connect("dbname='mydb' user='admin' password='admin' host='localhost' port='5432'")
        cur = conn.cursor()
        count_line=1
        count=1
        parsed = parser.from_file(pdf)
        text=parsed["content"]    
        text=text.split("\n")
        text=" ".join(text)
        po_date=re.search("PO Date.{20}",text).group().split(":")[1].split("Vendor")[0]
        po_no=re.search("PO#.{35}",text).group().split(":")[1].split("PO Date")[0]
        vendor_name=re.search("Hospital Name.*?Hos",text).group()
        vendor_name=vendor_name.replace("vendor_name","").replace("Hos","").split(" Name :")[1]

        try:
            text3=re.search("PO\s+Type :.*?Purchase",text).group().replace("PO Type","").replace(":","")
            text4=re.search("Hospital\s+Name.*?Hospital\s+Address",text).group()
            text5=path.split("/")[-1]
            text2=o+""+text4+""+text5+""+text3
            text2=text2.lower()
            mapping_list=["bellandur","chandigarh","noida","oar fertility op","hrbr ip","jayanagar ip","malleshwaram","oar ip","whitefield ip","shivajinagar pune","kalyani nagar pune","vashi","mumbai","noida ip","malleswaram op","chandigarh ip","gurgaon ip","oar fertility ip","jayanagar op","noida op","chandigarh op","gurgaon op","chennai ip","chennai op","bellandur ip","gurgoan sector 14"]
            for list_value in mapping_list:
                mapping_value=list_value.split(" ")
                if all(word in text2 for word in mapping_value):
                    mapping_word=list_value
                    mapping_dict={"bellandur":"C9 BELLANDUR IP","chandigarh":"C9 Chandigarh OP","noida":"C9 NOIDA IP","oar fertility op":"C9 FERTILITY OAR OP","hrbr ip":"C9 HRBR IP","jayanagar ip":"C9 JNR IP","malleshwaram":"C9 MWM IP","oar ip":"C9 OAR IP","whitefield ip":"C9 WHF IP",'shivajinagar pune': 'C9 SHIVAJINAGAR, PUNE', 'kalyani nagar pune': 'C9 KALYANINAGAR, PUNE', 'vashi': 'C9 VASHI', 'mumbai': 'C9 MALAD', 'noida ip': 'C9 NOIDA IP', 'malleswaram op': 'C9 MWM OP', 'chandigarh ip': 'C9 CHANDIGARH IP', 'gurgaon ip': 'C9 GURGAON IP', 'oar fertility ip': 'C9 FERTILITY OAR IP', 'jayanagar op': 'C9 JNR OP', 'noida op': 'C9 NOIDA OP', 'chandigarh op': 'C9 CHANDIGARH OP', 'gurgaon op': 'C9 GURGAON OP', 'chennai ip': 'C9 TNG IP', 'chennai op': 'C9 TNG OP', 'bellandur ip': 'C9 BELLANDUR IP', 'gurgoan sector 14': 'C9 GURGAON, SEC 14'}
            for key,value in mapping_dict.items():
                if mapping_word in key:
                    vendor_name=value
                    print(vendor_name)

        except:
            vendor_name="MOTHERHOOD DEMO"

        if check=='Normal':
            return vendor_name,po_no,po_date
        elif 'Line' in check:
            parsed = parser.from_file(pdf)
            text=parsed["content"]
            text1=text.split("\n")

            for num,x in enumerate(text1):                
                if re.search(f"^{count}\s+\S+",x):
                    value=re.search(f"{count}\s+",x).group()
                    text1[num]=x.replace(x,f"ITEM-{x}",1)
                    count=count+1
    
                elif str(x)==str(count):
                    text1[num]=x.replace(x,f"ITEM-{x}",count)
                    count=count+1
                    print("else")                        
            text=" ".join(text1)
            for main_data in text.split("ITEM"):
                if " ---------" in main_data:
                    main_data=main_data.split(" ---------")[0]
                if re.search("\d+\s+(EACH|BOTTLE|UNIT)\s+(EACH|BOTTLE|UNIT)",main_data):
                    if 'Check' in check:         
                        return len(text.split("ITEM"))-1
                    main_data=main_data.split(" PO Raised By")[0]
                    print("main_data",main_data)                
                    des=main_data.split(re.search("\d+\s+(EACH|BOTTLE|UNIT)\s+(EACH|BOTTLE|UNIT)",main_data).group())[0]
                    uom_quan=re.search("\d+\s+(EACH|BOTTLE|UNIT)\s+(EACH|BOTTLE|UNIT)",main_data).group()
                    remaing_data=main_data.split(uom_quan)[1]
                    #print(uom_quan)
                    remaing_data=remaing_data.split()
                    quantity=uom_quan.split()[0]
                    uom=uom_quan.split()[1]
                    print("remaing_data",remaing_data,"---------CODE IS ENTERED TO TIKA PART---------",len(remaing_data))                
                    total_amount=remaing_data[-3]
                    print(total_amount)                                     
                    if re.search("\d{8}",des):
                        des=des.split(re.search("\d{8}",des).group())[1]
                    if len(remaing_data)==11:
                        #des=main_data.split(uom_quan)[0]
                        rate,UMRP,TUMRP=extract_line_item(remaing_data,0,3,4)
                        if UMRP=='nan':
                            UMRP=remaing_data[4]
                            TUMRP=remaing_data[5]
                        #display(des,rate,uom,quantity,UMRP,TUMRP,discount_per)            
                    elif len(remaing_data)==12:
                        rate,UMRP,TUMRP=extract_line_item(remaing_data,0,3,4)
                        display(des,rate,uom,quantity,UMRP,TUMRP,discount_per) 
                        if UMRP=='nan':
                            UMRP=remaing_data[4]
                            TUMRP=remaing_data[5]
                        #display(des,rate,uom,quantity,UMRP,TUMRP,discount_per)

                    elif len(remaing_data)==9:
                        rate,UMRP,TUMRP=extract_line_item(remaing_data,0,3,4)
                        #display(des,rate,uom,quantity,UMRP,TUMRP,discount_per)

                    print(float(rate)*float(quantity),'---------------',total_amount)
                    if float(total_amount)==float(rate)*float(quantity):
                        discount_amount=0
                        discount_percent=0
                    else:
                        discount_amount=float(rate)*float(quantity)-float(total_amount)
                        discount_amount=round(discount_amount,2)
                        print("discount_amount",discount_amount)
                        discount_percent=round((discount_amount/(float(rate)*float(quantity)))*100)
                        print("discount_percent",discount_percent)
                        if discount_amount<1:
                            discount_amount=0                        
                    display(des,rate,uom,quantity,UMRP,TUMRP,discount_percent)
                    if round(float(quantity)*float(UMRP),2)==float(TUMRP):
                        print("**********fine*********"+str(count_line))                        
                        count_line=count_line+1                    
                    else:
                        exceptions_list.append("Issue in UMRP")
                    
                    if len(uom)>15 or re.search("\d+",uom):
                        exceptions_list.append("Issue in UOM") 


    #extract_normal_feild(r"C:\Users\acer\Desktop\office\kids\C9\kids3\New folder\kids_sample57.pdf","Line")


    def separator(data):
        data=data.strip()
        if data.count(".")==2 and re.search("\d+\s\d+",data)==None:        
            data=data.split(".")
            #print(data)
            final_data=data[0]+"."+data[1][:2]+" "+data[1][2:]+"."+data[2][:2]
            #print(final_data)
            return final_data
        if data.count(".")==3  and re.search("\s",data)==None:
            data=data.split(".")
            #print(data)
            final_data=data[0]+"."+data[1][:2]+" "+data[1][2:]+"."+data[2][:2]+" "+data[2][2:]+"."+data[3][:2]
            #print("",final_data)
            return final_data

        
      
        if len(data.split())==2:
           data1,data2=data.split()[0],data.split()[1]
           if data1.count('.')==2 and  data2.count('.')==1:
               data1=data1.split('.')
               print("....",data1)           
               data=data1[0]+"."+data1[1][:2]+" "+data1[1][2:]+"."+data1[2][:2]+" "+data2
               print("...........",data)
                     
            
        return data

            


    def extract(excel_file,pdf):
        conn=psycopg2.connect("dbname='mydb' user='admin' password='admin' host='localhost' port='5432'")
        cur = conn.cursor()
        wb=openpyxl.load_workbook("excel.xlsx")
        sh1=wb.active
        error=""
        exceptions_list=[]
        count=1
        f=excel_file
        print()
        print()
        print()
        print("excel file",f)
        print("@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@")
        vendor_name,po_no,po_date=extract_from_tika(pdf,"Normal")
        try:
            df=pd.read_csv(f)
        except Exception as e:
            error=str(e)
            
        if "No columns to parse from file" in error:
            f=f.replace("-pdf-page-1-tables.csv",".pdf")
            print(f)
            parsed = parser.from_file(pdf)
            text=parsed["content"]    
            text=text.split("\n")
            text=" ".join(text)
            print(text)

            line_item=re.search("Tax Value.*PO Raised By",text).group().replace("Amount Value Value","").replace("PO Raised By","")
            print(line_item)

            if re.search("\d+\s+(EACH|BOTTLE|UNIT)\s+(EACH|BOTTLE|UNIT)",line_item):
                uom_quan=re.search("\d+\s+(EACH|BOTTLE|UNIT)\s+(EACH|BOTTLE|UNIT)",line_item).group()            
                remaing_data=line_item.split(uom_quan)[1].split()
                print(remaing_data)
                quantity=uom_quan.split()[0]
                uom=uom_quan.split()[1]
                des=line_item.split(uom_quan)[0]
                discount_per=remaing_data[-2]
                total_amount=remaing_data[-3]
                if re.search("\d{8}",des):
                    des=des.split(re.search("\d{8}",des).group())[1]
                if len(remaing_data)==11:
                    #des=main_data.split(uom_quan)[0]
                    rate,UMRP,TUMRP=extract_line_item(remaing_data,0,3,4)
                    if UMRP=='nan':
                        UMRP=remaing_data[4]
                        TUMRP=remaing_data[5]
                    display(des,rate,uom,quantity,UMRP,TUMRP,discount_per)


                    if float(total_amount)==float(rate)*float(quantity):
                        discount_amount=0
                        discount_percent=0
                    else:
                        discount_amount=float(rate)*float(quantity)-float(total_amount)
                        discount_amount=round(discount_amount,2)
                        print("discount_amount",discount_amount)
                        discount_percent=round((discount_amount/(float(rate)*float(quantity)))*100)
                        print("discount_percent",discount_percent)
                        if (100-discount_percent)<discount_percent:
                            discount_percent=100-discount_percent
                        if discount_amount<1:
                            discount_amount=0
                    
                    disc_percent=discount_percent
                    foc,total_disc_percent,total_disc_amount,disc_amount='','','',''
                    sgst,igst,cgst='','',''
                    free=""
                    tax=""
                    qty=quantity
                    prate=rate
                    mrp=UMRP
                    total_disc_amount=discount_amount
                    time_stamp=""
                    pdf_file=pdf
                    time_stamp=str(datetime.datetime.today())
                    list_data = [vendor_name, po_no, po_date,des,qty,uom,mrp,foc,prate,disc_percent,disc_amount,total_disc_percent,total_disc_amount,sgst,igst,cgst,pdf_file,time_stamp]
                    sh1.append(list_data)                        
                    wb.save("excel.xlsx")
                    for a,b,c,d,e,f,g,h,i,j,k,l,m,n,o,p,q,r in [list_data]:
                        cur.execute("INSERT INTO vendor_api VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)",(a,b,c,d,e,f,g,h,i,j,k,l,m,n,o,p,q,r,"",'API Posting Pending','API Posting Pending'))
                        conn.commit()
                    conn.close()
            return "done"

     
        df.to_excel("test.xlsx",index=False)
        time.sleep(1)
        df=pd.read_excel("test.xlsx")
        for i in range(df.shape[0]):
            discount_percent=''
            des,rate,uom,quantity,UMRP,TUMRP="","","","","",""
            data=df.iloc[i,::].values.tolist()
            try:
                next_line=df.iloc[i+1,::].values.tolist()
                #print("next_line",next_line)
            except Exception as e:
                print(e)        
            print("---------------------------------------------------------",len(data),data)
            for n,y in enumerate(next_line):
                next_line[n]=str(y)

                
            def description(des1,des2):
                des=des1+" "+des2
                print(des)   
                return des
            if 'Table' in data :
                print("break-------------------------------------------------------------")
                break
            if 'Qty UOM' not in data and 'Purchase ' not in data and 'Rate ' not in data and 'Discounted Value ' not in data and len(data)>=7:
                for n,y in enumerate(data):
                    data[n]=str(y)
                print(data)
                separated_value=separator(data[5])
                data[5]=separated_value
                separated_value=separator(data[6])
                data[6]=separated_value

                try:
                    separated_value=separator(data[8])
                    data[8]=separated_value
                except:
                    pass
                try:
                    separated_value=separator(data[7])
                    data[7]=separated_value
                except:
                    pass
                separated_value=separator(data[3])
                data[3]=separated_value
                try:
                    separated_value=separator(data[9])
                    data[9]=separated_value
                except:
                    pass
                try:
                    separated_value=separator(data[10])
                    data[10]=separated_value
                except:
                    pass
                print("---------------------------------------------------------",len(data),"--------------------",data)
                main_data=" ".join(data).replace('nan','',1)
                print("=========================****************+++++++++++++++++++++++++++++++",main_data)
         

            
                
                if re.search("\d+\s+(EACH|BOTTLE|UNIT)\s+\d+.\d+|\d+\s+(EACH|BOTTLE|UNIT)\s+\d+\s+\d+\.\d+",main_data):
                    discount_per=""
                    discount_percent=''
                    #second format of pdf
                    print("second format of pdf")
                    uom_quan=re.search("\d+\s+(EACH|BOTTLE|UNIT)",main_data).group()
                    remaing_data=main_data.split(uom_quan)[1]
                    #print(uom_quan)
                    remaing_data=remaing_data.split()
                    quantity=uom_quan.split()[0]
                    uom=uom_quan.split()[1]
                    print("$$$$$$$$$$$$$$$$$$$$$$$remaing_data",remaing_data,"-------EXTRACTING DATA FRO AWS--------",len(remaing_data))                
                    total_amount=remaing_data[-1]
                    des=data[1].strip()
                    if next_line.count("nan")>=5:
                        des=des+" "+next_line[1]
                    if re.search("\d{8}",des):
                        des=des.split(re.search("\d{8}",des).group())[1]

                    
                    total_amount=remaing_data[-4]
                    print(total_amount)
                    
                    if len(remaing_data)==4:
                        discount_percent=0
                        rate,UMRP,TUMRP=extract_line_item(remaing_data,0,2,3)
                    elif len(remaing_data)==13:
                        print("-------EXTRACTING DATA FRO AWS 2nd format 13--------")
                        rate,UMRP,TUMRP=extract_line_item(remaing_data,0,1,3)
                    elif len(remaing_data)==12:
                        print("-------EXTRACTING DATA FRO AWS 2nd format 12--------")
                        rate,UMRP,TUMRP=extract_line_item(remaing_data,0,1,3)
                    elif len(remaing_data)==14:
                        print("-------EXTRACTING DATA FRO AWS 2nd format 14--------")
                        rate,UMRP,TUMRP=extract_line_item(remaing_data,0,1,3)
                    if str(TUMRP)=='nan':
                        TUMRP=remaing_data[4]
                    TUMRP=re.sub("\d{6}","",TUMRP)
                    
                    #display(des,rate,uom,quantity,UMRP,TUMRP,discount_percent)
                
                    if  discount_percent!=0:
                        if float(total_amount)==float(rate)*float(quantity):
                            discount_amount=0
                            discount_percent=0
                        else:
                            discount_amount=float(rate)*float(quantity)-float(total_amount)
                            discount_amount=round(discount_amount,2)
                            print("discount_amount",discount_amount)
                            discount_percent=round((discount_amount/(float(rate)*float(quantity)))*100)
                            print("discount_percent",discount_percent)
                            if (100-discount_percent)<discount_percent:
                                discount_percent=100-discount_percent
        
                            if discount_amount<1:
                                discount_amount=0
                            if "." in TUMRP:
                                TUMRP=TUMRP.split('.')
                                TUMRP=TUMRP[0]+"."+TUMRP[1][:2]
                        print(">>>>>>>>>>>>>>>>>>")
        
                        display(des,rate,uom,quantity,UMRP,TUMRP,discount_percent)

                        disc_percent=discount_percent
                        foc,total_disc_percent,total_disc_amount,disc_amount='','','',''
                        sgst,igst,cgst='','',''
                        free=""
                        tax=""
                        qty=quantity
                        prate=rate
                        mrp=UMRP
                        total_disc_amount=discount_amount
                        time_stamp=""
                        pdf_file=pdf
                        time_stamp=str(datetime.datetime.today())
                        list_data = [vendor_name, po_no, po_date,des,qty,uom,mrp,foc,prate,disc_percent,disc_amount,total_disc_percent,total_disc_amount,sgst,igst,cgst,pdf_file,time_stamp]
                        sh1.append(list_data)                        
                        wb.save("excel.xlsx")
                        for a,b,c,d,e,f,g,h,i,j,k,l,m,n,o,p,q,r in [list_data]:
                            cur.execute("INSERT INTO vendor_api VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)",(a,b,c,d,e,f,g,h,i,j,k,l,m,n,o,p,q,r,"",'API Posting Pending','API Posting Pending'))
                            conn.commit()
                       
                    continue 

                #print("=========================****************+++++++++++++++++++++++++++++++",main_data)
                final_check=checker2(main_data)                
                if final_check=="Tika":
                    extract_from_tika(pdf,"Line")
                    print("*********************************************************")
                    break                
                    
                elif  re.search("\d+\s+(EACH|BOTTLE|UNIT)\s+(EACH|BOTTLE|UNIT)",main_data):

                    uom_quan=re.search("\d+\s+(EACH|BOTTLE|UNIT)\s+(EACH|BOTTLE|UNIT)",main_data).group()
                    remaing_data=main_data.split(uom_quan)[1]
                    #print(uom_quan)
                    remaing_data=remaing_data.split()
                    quantity=uom_quan.split()[0]
                    uom=uom_quan.split()[1]
                    print("remaing_data",remaing_data,"-------EXTRACTING FROM AWS IST FORMAT--------",len(remaing_data))                
                    total_amount=remaing_data[-1]
                    des=data[1].strip()
                    discount_per=remaing_data[-2]
                    total_amount=remaing_data[-3]
                    print(des)
            
           
                    if len(des)==8 and re.search("\d{8}",des):
                        des=data[2]
                        if next_line.count("nan")>=8:
                            des=des+" "+next_line[2]
                    elif ("EACH" or 'BOTTLE' or 'UNIT') in des:
                        des=data[0]
                        if next_line.count("nan")>=7:
                            des=des+" "+next_line[0]
                    elif re.search("\d{8}\.\d",des) and len(des.strip())<12:                    
                        des=data[2]
                        if next_line.count("nan")>=8:
                            des=des+" "+next_line[2]                    
                    else:
                        if next_line.count("nan")>=8:
                            des=des+" "+next_line[1]

                    print(des)


                    def description_check(des):
                        regex='jdgfiebfjkbdfd'
                        if 'SURGICALS' in des.strip().split()[-1]:
                            des=des.strip().split()
                            des=" ".join(des[:len(des)-1])
                        if 'tab' in des[:6]:
                            des=des.split("tab")
                            des=" ".join(des[1:])
                            regex="\d+\stab"
                            
                        if '(ROMSONS)' in des[:11]:
                            des=des.split("(ROMSONS)")
                            des=" ".join(des[1:])
                            regex='(ROMSONS)'                                        
                        des2= next_line[1]                        
                        if re.search(regex,des2[:11]):
                            des2=re.search(regex,des2).group()
                            des=des+" "+des2
                            print(des)
                    
                    

                        print(des)
            
                        if re.search("\d{8}",des):
                            des1=des.split(re.search("\d{8}",des).group())[1]
                        else:
                            des1=des
                        des1=des1.lower()
                        des1= re.sub('\d+','',des1)
                        des1=re.sub('set','',des1)
                        des1=des1.replace(' ','')                
                        if len(des1)>8:
                            return des
                        else:
                            des=re.sub('\d{8}','',des)
                            return des

                        
                        
                    des=description_check(des)                

                    if re.search("\d{8}",des):
                        des=des.split(re.search("\d{8}",des).group())[1]                
                    if len(remaing_data)==11:
                        print("-------EXTRACTING FROM AWS IST FORMAT remaing_data11--------")
                        #des=main_data.split(uom_quan)[0]
                        rate,UMRP,TUMRP=extract_line_item(remaing_data,0,3,4)
                        if UMRP=='nan':
                            UMRP=remaing_data[4]
                            TUMRP=remaing_data[5]
                        #display(des,rate,uom,quantity,UMRP,TUMRP,discount_per)            
                    elif len(remaing_data)==12:
                        print("-------EXTRACTING FROM AWS IST FORMAT remaing_data12--------")
                        rate,UMRP,TUMRP=extract_line_item(remaing_data,0,3,4)
                        display(des,rate,uom,quantity,UMRP,TUMRP,discount_per) 
                        if UMRP=='nan':
                            UMRP=remaing_data[4]
                            TUMRP=remaing_data[5]
                        #display(des,rate,uom,quantity,UMRP,TUMRP,discount_per)
                    elif len(remaing_data)==10:
                        print("-------EXTRACTING FROM AWS IST FORMAT remaing_data10--------")
                        rate,UMRP,TUMRP=extract_line_item(remaing_data,0,3,4)
                        #display(des,rate,uom,quantity,UMRP,TUMRP,discount_per)
                    elif len(remaing_data)==9:
                        print("-------EXTRACTING FROM AWS IST FORMAT remaing_data9--------")
                        rate,UMRP,TUMRP=extract_line_item(remaing_data,0,3,4)
                        if float(UMRP)>float(TUMRP):
                            UMRP=0
                            TUMRP=0                        
                      #display(des,rate,uom,quantity,UMRP,TUMRP,discount_per)
                    elif len(remaing_data)==8:
                        print("-------EXTRACTING FROM AWS IST FORMAT remaing_data8--------")
                        rate,UMRP,TUMRP=extract_line_item(remaing_data,0,3,4)
                        #display(des,rate,uom,quantity,UMRP,TUMRP,discount_per) 
                    else:
                        print("###############################",main_data)            
                    #display(des,rate,uom,quantity,UMRP,TUMRP,discount_percent)


                    if float(total_amount)==float(rate)*float(quantity):
                        discount_amount=0
                        discount_percent=0
                    else:
                        discount_amount=float(rate)*float(quantity)-float(total_amount)
                        discount_amount=round(discount_amount,2)
                        print("discount_amount",discount_amount)
                        discount_percent=round((discount_amount/(float(rate)*float(quantity)))*100)
                        print("discount_percent",discount_percent)
                        if (100-discount_percent)<discount_percent:
                            discount_percent=100-discount_percent                        
                        if discount_amount<1:
                            discount_amount=0
                    if discount_percent ==6:
                        discount_percent=5
                        
                    #print("xxxxxxxxxxxxxxxxxx",total_amount)      
                    display(des,rate,uom,quantity,UMRP,TUMRP,discount_percent)
                    if round(float(quantity)*float(UMRP),2)==float(TUMRP):# and math.isclose(float(round(float(rate)*float(quantity)- float(discount_amount))),float(total_amount), abs_tol=10):
                        print("**********fine*********"+str(count))                        
                        count=count+1                    
                    else:
                        exceptions_list.append("Issue in UMRP")                
                    if len(uom)>15 or re.search("\d+",uom):
                        exceptions_list.append("Issue in UOM")
                    disc_percent=discount_percent
                    foc,total_disc_percent,total_disc_amount,disc_amount='','','',''
                    sgst,igst,cgst='','',''
                    free=""
                    tax=""
                    qty=quantity
                    prate=rate
                    mrp=UMRP
                    total_disc_amount=discount_amount
                    time_stamp=""
                    pdf_file=pdf
                    time_stamp=str(datetime.datetime.today())
                    list_data = [vendor_name, po_no, po_date,des,qty,uom,mrp,foc,prate,disc_percent,disc_amount,total_disc_percent,total_disc_amount,sgst,igst,cgst,pdf_file,time_stamp]
                    sh1.append(list_data)                        
                    wb.save("excel.xlsx")
                    
                    for a,b,c,d,e,f,g,h,i,j,k,l,m,n,o,p,q,r in [list_data]:
                        cur.execute("INSERT INTO vendor_api VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)",(a,b,c,d,e,f,g,h,i,j,k,l,m,n,o,p,q,r,"",'API Posting Pending','API Posting Pending'))
                        conn.commit()
##                    conn.close()


    input_pdf_path=r"/home/ubuntu/Desktop/Automation/FINAL_CODES/FINAL_CODES"
    for data in os.listdir(input_pdf_path):
           data2= (input_pdf_path+"/%s")%data
           csv1=data2
           name=name.replace(".pdf","")
          
           
        
           if "page-1-tables"  in data and '.csv' in data2 and name in data2:            
                #print(data2)
                csv2=data2.replace("-1-","-2-")
                #print(csv1,"------------",csv2)
                pdf=csv1.replace("-pdf-page-1-tables.csv",".pdf")
                line_item_from_tika=extract_from_tika(pdf,'Line-Check')
                #print("line_item_from_tika",line_item_from_tika)           
                path1=csv1
                path2=csv2
                try:
                    combine_csv(path1,path2)
                except:
                    line_item_from_aws='NUN'
                    
                line_item_from_aws=count_line("final.csv")      
                print("count_line:----------------------",line_item_from_tika,line_item_from_aws)
        
                if line_item_from_tika!=line_item_from_aws and line_item_from_tika!=None:
                    extract_from_tika(pdf,'Line')
                   
                else:
                    extract("final.csv",pdf)


##    for data in os.listdir(input_pdf_path):
##           data2= (input_pdf_path+"/%s")%data
##        
##           if ( 'py' ) not in data2 and  '__pycache__' not in data2 and name in data2 :
##               print("xxxxxxxxxxx",data2)
##               try:
##                   shutil.move(data2, '/home/ubuntu/Desktop/Automation/kids_data_backup')
##               except:
##                   pass

##extract(r"/home/ubuntu/Desktop/Automation/FINAL_CODES/FINAL_CODES/FETCHED_PDF/purchase_order_311492.pdf","")
